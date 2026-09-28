"""解析测试问题反馈表。

识别含「序号」和「故障概述」列的工作表，跳过填写说明、空模板等其他表。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from pathlib import Path


class IssueSheetError(Exception):
    pass


_VERSION_RE = re.compile(r"V\d+(?:\.\d+){2,}", re.IGNORECASE)
_FULL_DATE_RE = re.compile(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})")
_MONTH_DAY_RE = re.compile(r"(\d{1,2})月(\d{1,2})日")
_SOFTWARE_RE = re.compile(r"软件版本\s*[:：]\s*([^\s/]+)")
_MCU_RE = re.compile(r"MCU\s*[:：]\s*([^\s/]+)", re.IGNORECASE)
_NAMED_VERSION_RE = re.compile(r"([A-Za-z][A-Za-z0-9_-]{1,24})\s*[:：]\s*(V[0-9][^\s/]*)")

_PRIORITY_BY_LEVEL = {
    "4veryhigh": "紧急",
    "veryhigh": "紧急",
    "非常高": "紧急",
    "紧急": "紧急",
    "3high": "高",
    "high": "高",
    "高": "高",
    "2medium": "普通",
    "medium": "普通",
    "中": "普通",
    "普通": "普通",
    "1low": "低",
    "low": "低",
    "低": "低",
    "5immediate": "立刻",
    "immediate": "立刻",
    "立刻": "立刻",
}


@dataclass
class ParsedIssue:
    row_key: str
    sheet: str
    excel_row: int
    seq: str
    seq_note: str
    found_on: str
    found_raw: str
    module: str
    level: str
    priority_name: str
    tracker_name: str
    subject: str
    description: str


@dataclass
class ParsedWorkbook:
    filename: str
    title: str = ""
    software_version: str = ""
    mcu: str = ""
    extra_versions: list[tuple[str, str]] = field(default_factory=list)
    sample: str = ""
    suggested_version: str = ""
    issues: list[ParsedIssue] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    used_sheets: list[str] = field(default_factory=list)


def parse_issue_workbook(filename: str, content: bytes, *, year: int | None = None) -> ParsedWorkbook:
    sheets = read_workbook_sheets(filename, content)
    if not sheets:
        raise IssueSheetError("表格里没有工作表")
    resolved_year = year or date.today().year
    parsed = ParsedWorkbook(filename=Path(filename).name)
    for sheet_name, rows in sheets:
        issue_rows, context, row_warnings = _parse_sheet(sheet_name, rows, parsed.filename, resolved_year)
        if issue_rows is None:
            parsed.skipped_sheets.append(sheet_name)
            parsed.warnings.append(f"已跳过工作表「{sheet_name}」：没有「序号」和「故障概述」列")
            continue
        parsed.used_sheets.append(sheet_name)
        parsed.warnings.extend(row_warnings)
        if not parsed.title:
            parsed.title = context["title"]
            parsed.software_version = context["software_version"]
            parsed.mcu = context["mcu"]
            parsed.extra_versions = context["extra_versions"]
            parsed.sample = context["sample"]
        parsed.issues.extend(issue_rows)
    if not parsed.issues:
        raise IssueSheetError("没有解析到问题。请确认表格中有「序号」和「故障概述」列")
    parsed.suggested_version = _suggest_version(parsed)
    duplicate_seqs = sorted({item.seq for item in parsed.issues if item.seq_note})
    if duplicate_seqs:
        parsed.warnings.append("表中序号重复，已按多条问题保留：" + "、".join(duplicate_seqs))
    return parsed


def read_workbook_sheets(filename: str, content: bytes) -> list[tuple[str, list[list[str]]]]:
    suffix = Path(filename or "").suffix.lower()
    readers = (_read_xls, _read_xlsx) if suffix == ".xls" else (_read_xlsx, _read_xls)
    errors: list[str] = []
    for reader in readers:
        try:
            return reader(content)
        except IssueSheetError:
            raise
        except Exception as exc:
            errors.append(str(exc))
    detail = "；".join(item for item in errors if item) or "文件不是可读取的 Excel 表格"
    raise IssueSheetError(f"无法读取表格：{detail}")


def priority_name_for_level(level: str) -> str:
    key = re.sub(r"[\s_\-]+", "", level or "").casefold()
    return _PRIORITY_BY_LEVEL.get(key, "")


def tracker_name_for_summary(summary: str) -> str:
    return "功能" if "建议" in (summary or "") else "错误"


def _read_xlsx(content: bytes) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    try:
        sheets: list[tuple[str, list[list[str]]]] = []
        for worksheet in workbook.worksheets:
            rows = [
                [_cell_text(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            sheets.append((worksheet.title, rows))
        return sheets
    finally:
        workbook.close()


def _read_xls(content: bytes) -> list[tuple[str, list[list[str]]]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheets: list[tuple[str, list[list[str]]]] = []
    for worksheet in book.sheets():
        rows: list[list[str]] = []
        for row_index in range(worksheet.nrows):
            values: list[str] = []
            for col_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, col_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    parsed = xlrd.xldate_as_datetime(cell.value, book.datemode)
                    values.append(parsed.date().isoformat())
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    values.append(str(int(cell.value)))
                else:
                    values.append(_cell_text(cell.value))
            rows.append(values)
        sheets.append((worksheet.name, rows))
    return sheets


def _cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return "" if not isinstance(value, bool) else ("1" if value else "")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.replace("\u00a0", " ").replace("\u3000", " ").strip()


def _parse_sheet(
    sheet_name: str,
    rows: list[list[str]],
    filename: str,
    year: int,
) -> tuple[list[ParsedIssue] | None, dict[str, object], list[str]]:
    header_index = _find_header_row(rows)
    if header_index is None:
        return None, {}, []
    header = [_compact(cell) for cell in rows[header_index]]
    columns = {
        "seq": _find_exact(header, "序号"),
        "found": _find_contains(header, "发现日期"),
        "module": _find_contains(header, "测试项目", "所属模块"),
        "sample_no": _find_contains(header, "故障样品"),
        "summary": _find_contains(header, "故障概述"),
        "detail": _find_contains(header, "详细说明", "描述"),
        "level": _find_contains(header, "故障等级", "严重度"),
        "status": _find_contains(header, "当前现状"),
        "tester": _find_contains(header, "测试人员", "发现人"),
        "solution": _find_contains(header, "分析或解决方案"),
        "owner": _find_contains(header, "责任人"),
    }
    context = _sheet_context(rows[:header_index])
    warnings: list[str] = []
    if columns["level"] is None:
        warnings.append(f"工作表「{sheet_name}」没有故障等级列，未填写等级的问题按普通处理")
    issues: list[ParsedIssue] = []
    seen_seq: dict[str, int] = {}
    for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        summary = _at(row, columns["summary"])
        seq = _at(row, columns["seq"])
        if not summary and not seq:
            continue
        if not summary:
            warnings.append(f"工作表「{sheet_name}」第 {offset} 行没有故障概述，已跳过")
            continue
        seq_note = ""
        if seq:
            seen_seq[seq] = seen_seq.get(seq, 0) + 1
            if seen_seq[seq] > 1:
                seq_note = "（表中序号重复）"
        found_raw = _at(row, columns["found"])
        found_on = _parse_found_on(found_raw, year)
        if found_raw and not found_on:
            warnings.append(f"工作表「{sheet_name}」第 {offset} 行日期无法识别：{found_raw}")
        level = _at(row, columns["level"])
        priority_name = priority_name_for_level(level) if level else "普通"
        if level and not priority_name:
            warnings.append(f"工作表「{sheet_name}」第 {offset} 行故障等级无法识别：{level}")
        detail = _at(row, columns["detail"]) or summary
        subject = _subject(summary)
        description = _description(
            filename=filename,
            sheet_name=sheet_name,
            seq=seq or "（无序号）",
            seq_note=seq_note,
            found_on=found_on,
            found_raw=found_raw,
            module=_at(row, columns["module"]),
            sample_no=_at(row, columns["sample_no"]),
            context=context,
            level=level,
            summary=summary,
            detail=detail,
            status=_at(row, columns["status"]),
            tester=_at(row, columns["tester"]),
            solution=_at(row, columns["solution"]),
            owner=_at(row, columns["owner"]),
        )
        issues.append(
            ParsedIssue(
                row_key=f"{sheet_name}\t{offset}",
                sheet=sheet_name,
                excel_row=offset,
                seq=seq,
                seq_note=seq_note,
                found_on=found_on,
                found_raw=found_raw,
                module=_at(row, columns["module"]),
                level=level,
                priority_name=priority_name,
                tracker_name=tracker_name_for_summary(summary),
                subject=subject,
                description=description,
            )
        )
    return issues, context, warnings


def _find_header_row(rows: list[list[str]]) -> int | None:
    for index, row in enumerate(rows):
        compact = [_compact(cell) for cell in row]
        if "序号" in compact and any(cell.startswith("故障概述") for cell in compact):
            return index
    return None


def _sheet_context(rows: list[list[str]]) -> dict[str, object]:
    cells = [cell for row in rows for cell in row if cell]
    blob = "\n".join(cells)
    software = _SOFTWARE_RE.search(blob)
    mcu = _MCU_RE.search(blob)
    extra: list[tuple[str, str]] = []
    seen: set[str] = set()
    for match in _NAMED_VERSION_RE.finditer(blob):
        name = match.group(1)
        if name.lower() == "mcu" or name in seen:
            continue
        seen.add(name)
        extra.append((name, match.group(2)))
    sample = next((cell for cell in cells if "样机" in cell and "测试样品状态" not in cell), "")
    return {
        "title": cells[0] if cells else "",
        "software_version": software.group(1) if software else "",
        "mcu": mcu.group(1) if mcu else "",
        "extra_versions": extra,
        "sample": sample,
    }


def _description(
    *,
    filename: str,
    sheet_name: str,
    seq: str,
    seq_note: str,
    found_on: str,
    found_raw: str,
    module: str,
    sample_no: str,
    context: dict[str, object],
    level: str,
    summary: str,
    detail: str,
    status: str,
    tester: str,
    solution: str,
    owner: str,
) -> str:
    extra_versions = context.get("extra_versions") or []
    lines = [
        f"来源：{filename} / {sheet_name}",
        f"序号：{seq}{seq_note}",
        f"发现日期：{found_on or found_raw or '（未填写）'}",
        f"测试项目或模块：{module or '（空）'}",
    ]
    if sample_no:
        lines.append(f"故障样品编号：{sample_no}")
    lines.extend(
        [
            f"软件版本：{context.get('software_version') or '（未填写）'}",
            f"MCU：{context.get('mcu') or '（未填写）'}",
        ]
    )
    for name, value in extra_versions:
        lines.append(f"{name}：{value}")
    lines.extend(
        [
            f"样品：{context.get('sample') or '（未填写）'}",
            f"故障等级：{level or '（未填写）'}",
            "",
            "故障概述：",
            summary,
            "",
            "故障及风险详细说明：",
            detail,
        ]
    )
    for label, value in (("当前现状", status), ("测试人员", tester), ("分析或解决方案", solution), ("表中责任人", owner)):
        if value:
            lines.extend(["", f"{label}：", value])
    return "\n".join(lines)


def _suggest_version(parsed: ParsedWorkbook) -> str:
    for candidate in (
        parsed.mcu,
        parsed.filename,
        parsed.software_version,
        *(value for _name, value in parsed.extra_versions),
    ):
        match = _VERSION_RE.search(candidate or "")
        if match:
            return match.group(0)
    return ""


def _parse_found_on(raw: str, year: int) -> str:
    text = (raw or "").strip()
    if not text:
        return ""
    full = _FULL_DATE_RE.fullmatch(text)
    if full:
        return _safe_date(int(full.group(1)), int(full.group(2)), int(full.group(3)))
    month_day = _MONTH_DAY_RE.fullmatch(text)
    if month_day:
        return _safe_date(year, int(month_day.group(1)), int(month_day.group(2)))
    return ""


def _safe_date(year: int, month: int, day: int) -> str:
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return ""


def _subject(summary: str) -> str:
    text = re.sub(r"\s+", " ", summary).strip()
    if len(text) <= 250:
        return text
    return text[:247].rstrip() + "..."


def _compact(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def _find_exact(header: list[str], expected: str) -> int | None:
    for index, cell in enumerate(header):
        if cell == expected:
            return index
    return None


def _find_contains(header: list[str], *needles: str) -> int | None:
    for needle in needles:
        for index, cell in enumerate(header):
            if needle in cell:
                return index
    return None


def _at(row: list[str], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return row[index].strip()
